"""Figure extraction. Pulls diagrams and images out of PDF pages and PPTX slides,
saves each one to disk once, and records it in the `figures` table. Gemini then
describes each figure (see kb.process_figures), and the description becomes a
searchable, citable knowledge-base unit with the same page / slide location."""
import hashlib
from pathlib import Path

from .store import connect

FIG_DIR = Path(__file__).resolve().parent.parent / "uploads" / "figures"
MIN_BYTES = 5000          # skip tiny logos, bullets and icons
MAX_PER_FILE = 30         # keeps Gemini usage bounded for one upload
VECTOR_DRAWINGS = 15      # pages with at least this many vector shapes are rendered
ALLOWED = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg", "webp": "image/webp"}


def _pdf_figures(path):
    figs = []
    try:
        import fitz
    except ImportError:
        fitz = None

    if fitz is not None:
        with fitz.open(path) as pdf:
            for i, page in enumerate(pdf, start=1):
                for img in page.get_images(full=True):
                    info = pdf.extract_image(img[0])
                    ext = info.get("ext", "").lower()
                    if ext in ALLOWED and len(info["image"]) >= MIN_BYTES:
                        figs.append({"location": f"Page {i}", "section": None,
                                     "mime": ALLOWED[ext], "ext": ext, "bytes": info["image"]})
                if len(page.get_drawings()) >= VECTOR_DRAWINGS:
                    png = page.get_pixmap(dpi=120).tobytes("png")
                    figs.append({"location": f"Page {i} (diagram)", "section": None,
                                 "mime": "image/png", "ext": "png", "bytes": png})
    else:  # pypdf fallback: embedded images only
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        for i, page in enumerate(reader.pages, start=1):
            try:
                images = list(page.images)
            except Exception:
                images = []
            for img in images:
                ext = Path(img.name).suffix.lower().lstrip(".")
                if ext in ALLOWED and len(img.data) >= MIN_BYTES:
                    figs.append({"location": f"Page {i}", "section": None,
                                 "mime": ALLOWED[ext], "ext": ext, "bytes": img.data})
    return figs


def _pptx_figures(path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    figs = []
    prs = Presentation(path)
    for i, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None else None
        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            img = shape.image
            ext = (img.ext or "").lower()
            if ext in ALLOWED and len(img.blob) >= MIN_BYTES:
                figs.append({"location": f"Slide {i}", "section": title,
                             "mime": ALLOWED[ext], "ext": ext, "bytes": img.blob})
    return figs


def extract(path):
    """Returns the figures found in one file (no saving, no Gemini call)."""
    suffix = Path(path).suffix.lower()
    if suffix == ".pdf":
        figs = _pdf_figures(path)
    elif suffix == ".pptx":
        figs = _pptx_figures(path)
    else:
        figs = []
    return figs[:MAX_PER_FILE]


def save_pending(doc_name, figs):
    """Writes each figure to uploads/figures and records it as pending. Returns count."""
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    count = 0
    with connect() as con:
        for f in figs:
            digest = hashlib.sha1(f["bytes"]).hexdigest()[:16]
            path = FIG_DIR / f"{digest}.{f['ext']}"
            if not path.exists():
                path.write_bytes(f["bytes"])
            exists = con.execute("SELECT id FROM figures WHERE path=?", (str(path),)).fetchone()
            if exists:
                continue
            con.execute(
                "INSERT INTO figures(source,location,section,mime,path,status) VALUES(?,?,?,?,?,'pending')",
                (doc_name, f["location"], f["section"], f["mime"], str(path)))
            count += 1
    return count


def pending_rows():
    with connect() as con:
        return [dict(r) for r in con.execute("SELECT * FROM figures WHERE status='pending' ORDER BY id")]


def counts():
    with connect() as con:
        rows = con.execute("SELECT status, COUNT(*) AS n FROM figures GROUP BY status").fetchall()
    out = {"pending": 0, "done": 0}
    for r in rows:
        out[r["status"]] = r["n"]
    return out


def mark_done(row_id, description, chunk_id):
    with connect() as con:
        con.execute("UPDATE figures SET status='done', description=?, chunk_id=? WHERE id=?",
                    (description, chunk_id, row_id))


def figure_for_chunk(chunk_id):
    with connect() as con:
        row = con.execute("SELECT * FROM figures WHERE chunk_id=?", (chunk_id,)).fetchone()
        return dict(row) if row else None


def for_location(source, location):
    """All figures on one page or slide, e.g. 'Page 3' also matches 'Page 3 (diagram)'."""
    with connect() as con:
        rows = con.execute("SELECT * FROM figures WHERE source=? ORDER BY id", (source,)).fetchall()
    out = []
    for r in rows:
        loc = r["location"]
        if loc == location or loc.startswith(location + " ("):
            out.append(dict(r))
    return out


def all_figures():
    with connect() as con:
        return [dict(r) for r in con.execute("SELECT * FROM figures ORDER BY id")]


def get_figure(fig_id):
    with connect() as con:
        row = con.execute("SELECT * FROM figures WHERE id=?", (fig_id,)).fetchone()
        return dict(row) if row else None
