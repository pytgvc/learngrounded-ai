"""Generates the demo course in sample_data/: a textbook PDF, a slide deck (PPTX)
and a lecture caption file (VTT). Run once: python make_sample_data.py"""
from pathlib import Path

from pptx import Presentation

try:
    import fitz
except ImportError:
    fitz = None
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent / "sample_data"
OUT.mkdir(exist_ok=True)

TEXTBOOK = [
    ("Chapter 3: Photosynthesis", [
        "Photosynthesis is the process by which green plants and some bacteria convert light energy into chemical energy. "
        "The process takes place mainly in the chloroplasts, which contain the green pigment chlorophyll. "
        "Chlorophyll absorbs light most strongly in the blue and red parts of the visible spectrum and reflects green light.",
        "The overall reaction uses six molecules of carbon dioxide and six molecules of water to produce one molecule of glucose and six molecules of oxygen. "
        "Light energy drives the reaction, so the process is effectively powered by sunlight.",
        "Photosynthesis has two main stages. The light-dependent reactions occur in the thylakoid membranes and split water to release oxygen. "
        "The Calvin cycle occurs in the stroma and uses ATP and NADPH from the light reactions to fix carbon dioxide into sugar.",
    ]),
    ("Chapter 4: Cellular Respiration", [
        "Cellular respiration releases the chemical energy stored in glucose and makes it available to the cell as ATP. "
        "Most of this process takes place in the mitochondria, which are often described as the powerhouse of the cell.",
        "Glycolysis splits one molecule of glucose into two molecules of pyruvate and produces a net gain of 2 ATP molecules. "
        "The Krebs cycle and the electron transport chain then extract much more energy, producing up to 36 ATP per glucose molecule under typical conditions.",
    ]),
    ("Chapter 5: The Water Cycle", [
        "The water cycle describes how water moves continuously between the oceans, the atmosphere, and the land. "
        "Evaporation from oceans and lakes transfers water vapor into the air, where it cools and condenses to form clouds.",
        "Precipitation returns water to the surface as rain, snow, or hail. Runoff carries water into rivers, and groundwater infiltration stores it underground for long periods.",
    ]),
]

SLIDES = [
    ("Cell Structure", ["The cell membrane controls what enters and leaves the cell.",
                        "The nucleus contains the genetic material DNA and is surrounded by a double membrane.",
                        "Ribosomes build proteins and can be free in the cytoplasm or attached to the endoplasmic reticulum."]),
    ("Mitosis", ["Mitosis produces two genetically identical daughter cells from one parent cell.",
                 "Mitosis has four main phases: prophase, metaphase, anaphase, and telophase.",
                 "Chromosomes line up at the cell equator during metaphase."]),
    ("Meiosis", ["Meiosis produces four haploid gametes, each with half the chromosome number of the parent cell.",
                 "Crossing over during prophase I shuffles genetic material and increases genetic variation."]),
]

VTT = [
    ("00:00:00.000", "00:00:08.000", "Welcome back. Today we look at the light-dependent reactions of photosynthesis."),
    ("00:00:08.000", "00:00:17.000", "These reactions happen in the thylakoid membranes inside the chloroplast."),
    ("00:00:17.000", "00:00:26.000", "When a photon hits chlorophyll, an electron gets excited and moves along an electron transport chain."),
    ("00:00:31.000", "00:00:40.000", "Water is split in a process called photolysis, which releases oxygen as a by-product."),
    ("00:00:40.000", "00:00:50.000", "The energy captured here is stored in ATP and NADPH for the Calvin cycle."),
    ("00:01:05.000", "00:01:15.000", "Next lecture we move on to the Calvin cycle and carbon fixation in the stroma."),
    ("00:01:15.000", "00:01:25.000", "Remember that the Calvin cycle needs ATP and NADPH from the light reactions to work."),
]


def make_pdf(path):
    """Writes a multi-page textbook PDF. Uses PyMuPDF if present, otherwise matplotlib."""
    import textwrap
    if fitz is not None:
        doc = fitz.open()
        for title, paras in TEXTBOOK:
            page = doc.new_page()
            page.insert_text((50, 60), title, fontsize=18, fontname="helv")
            y = 96
            for para in paras:
                page.insert_textbox(fitz.Rect(50, y, 545, 800), para, fontsize=11, fontname="helv")
                y += 110
            page.insert_text((50, 820), f"Page {len(doc)}", fontsize=9, fontname="helv")
        doc.save(path)
        return
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    with PdfPages(path) as pdf:
        for n, (title, paras) in enumerate(TEXTBOOK, start=1):
            fig = plt.figure(figsize=(8.5, 11))
            fig.text(0.08, 0.93, title, fontsize=18, weight="bold")
            y = 0.87
            for para in paras:
                wrapped = "\n".join(textwrap.wrap(para, 95))
                fig.text(0.08, y, wrapped, fontsize=10.5, va="top")
                y -= 0.035 * (wrapped.count("\n") + 3)
            fig.text(0.08, 0.03, f"Page {n}", fontsize=9)
            pdf.savefig(fig)
            plt.close(fig)


def make_diagram_png():
    """A labelled flow diagram of mitosis phases, embedded in the slide deck."""
    import io
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch
    fig, ax = plt.subplots(figsize=(9, 3.2))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 3)
    ax.axis("off")
    phases = ["Prophase", "Metaphase", "Anaphase", "Telophase"]
    for i, name in enumerate(phases):
        x = 0.4 + i * 2.4
        ax.add_patch(FancyBboxPatch((x, 1), 1.9, 1.1, boxstyle="round,pad=0.05",
                                    fc="#eef3ff", ec="#2f5bea", lw=2))
        ax.text(x + 0.95, 1.55, name, ha="center", va="center", fontsize=12, weight="bold")
        if i < len(phases) - 1:
            ax.annotate("", xy=(x + 2.3, 1.55), xytext=(x + 2.0, 1.55),
                        arrowprops=dict(arrowstyle="->", lw=2, color="#2f5bea"))
    ax.text(5, 2.5, "Figure: the four phases of mitosis, in order", ha="center", fontsize=11)
    ax.text(5, 0.4, "Chromosomes line up at the equator in metaphase", ha="center", fontsize=10)
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=120, bbox_inches="tight")
    plt.close(fig)
    return buf.getvalue()


def make_pptx(path):
    import io
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)
    for title, bullets in SLIDES:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.text = bullets[0]
        for b in bullets[1:]:
            body.add_paragraph().text = b
        for p in body.paragraphs:
            p.font.size = Pt(20)
        if title == "Mitosis":
            slide.shapes.add_picture(io.BytesIO(make_diagram_png()),
                                     Inches(0.8), Inches(4.6), width=Inches(11))
    prs.save(path)


def make_vtt(path):
    lines = ["WEBVTT", ""]
    for start, end, text in VTT:
        lines += [f"{start} --> {end}", text, ""]
    path.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    make_pdf(OUT / "Biology_Textbook.pdf")
    make_pptx(OUT / "Cell_and_Division_Slides.pptx")
    make_vtt(OUT / "Lecture_Photosynthesis.vtt")
    print("Sample course written to", OUT)
